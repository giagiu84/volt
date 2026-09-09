# -*- coding: utf-8 -*-
"""Genera VOLT_Bibbia.docx da STORIA.md.

La Bibbia si scrive in un posto solo: STORIA.md.
Il Word e' una stampa di quel file, per chi legge senza terminale.
Non si modifica il .docx a mano: si modifica l'.md e si rilancia questo.

    python strumenti/bibbia_docx.py
"""
import io, os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RADICE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SORGENTE = os.path.join(RADICE, 'STORIA.md')
USCITA = os.path.join(RADICE, 'VOLT_Bibbia.docx')

INCHIOSTRO = RGBColor(0x24, 0x1A, 0x4D)
BLU = RGBColor(0x22, 0x76, 0xC8)
VIOLA = RGBColor(0x7A, 0x3F, 0xD6)
GRIGIO = RGBColor(0x5B, 0x54, 0x77)


def sfondo(cella, esa):
    """Colore di riempimento di una cella: python-docx non lo espone."""
    tcPr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), esa)
    tcPr.append(shd)


def bordo_sinistro(par, esa):
    """Barretta colorata a sinistra: serve per le citazioni."""
    pPr = par._p.get_or_add_pPr()
    bordi = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '10')
    left.set(qn('w:color'), esa)
    bordi.append(left)
    pPr.append(bordi)


TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[\[.+?\]\])')


def scrivi(par, testo, base=None):
    """Porta in Word il grassetto, il corsivo e il codice del markdown."""
    for pezzo in TOKEN.split(testo):
        if not pezzo:
            continue
        if pezzo.startswith('**') and pezzo.endswith('**'):
            r = par.add_run(pezzo[2:-2]); r.bold = True
        elif pezzo.startswith('`') and pezzo.endswith('`'):
            r = par.add_run(pezzo[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            r.font.color.rgb = VIOLA
        elif pezzo.startswith('*') and pezzo.endswith('*'):
            r = par.add_run(pezzo[1:-1]); r.italic = True
        elif pezzo.startswith('[[') and pezzo.endswith(']]'):
            r = par.add_run(pezzo[2:-2]); r.italic = True
        else:
            r = par.add_run(pezzo)
        if base:
            r.font.color.rgb = base
    return par


def stili(doc):
    n = doc.styles['Normal']
    n.font.name = 'Calibri'
    n.font.size = Pt(11)
    n.paragraph_format.space_after = Pt(7)
    n.paragraph_format.line_spacing = 1.12
    for nome, dim, col in (('Heading 1', 19, INCHIOSTRO), ('Heading 2', 14.5, BLU),
                           ('Heading 3', 12, VIOLA)):
        s = doc.styles[nome]
        s.font.name = 'Calibri'
        s.font.size = Pt(dim)
        s.font.bold = True
        s.font.color.rgb = col
        s.paragraph_format.space_before = Pt(16 if nome == 'Heading 1' else 11)
        s.paragraph_format.space_after = Pt(5)


def copertina(doc, sottotitolo, data):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run('VOLT'); r.bold = True; r.font.size = Pt(64); r.font.color.rgb = INCHIOSTRO

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('BIBBIA NARRATIVA'); r.bold = True; r.font.size = Pt(20)
    r.font.color.rgb = BLU

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sottotitolo); r.font.size = Pt(12); r.font.color.rgb = GRIGIO

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run('NON SI NASCE VOLT. LO SI DIVENTA.')
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = VIOLA

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run(data); r.font.size = Pt(10); r.font.color.rgb = GRIGIO

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Generato da STORIA.md — non modificare questo file a mano:\n'
                  'si scrive nell\'.md e si rilancia strumenti/bibbia_docx.py')
    r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = GRIGIO
    doc.add_page_break()


def tabella(doc, righe):
    intestazione = [c.strip() for c in righe[0].strip('|').split('|')]
    corpo = [[c.strip() for c in r.strip('|').split('|')] for r in righe[2:]]
    t = doc.add_table(rows=1, cols=len(intestazione))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, testo in enumerate(intestazione):
        cel = t.rows[0].cells[i]
        cel.text = ''
        p = cel.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(re.sub(r'\*\*', '', testo)); r.bold = True; r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sfondo(cel, '3B4A7A')
    for k, riga in enumerate(corpo):
        celle = t.add_row().cells
        for i, testo in enumerate(riga[:len(intestazione)]):
            celle[i].text = ''
            p = celle[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            scrivi(p, testo)
            for r in p.runs:
                r.font.size = Pt(10)
            if k % 2:
                sfondo(celle[i], 'F1EFF8')
    doc.add_paragraph()


def converti():
    testo = io.open(SORGENTE, encoding='utf-8').read().split('\n')
    doc = Document()
    for sez in doc.sections:
        sez.top_margin = sez.bottom_margin = Cm(2.2)
        sez.left_margin = sez.right_margin = Cm(2.4)
    stili(doc)

    sottotitolo, data, sottotitolo_riga = '', '', ''
    for riga in testo[:25]:
        if riga.startswith('### ') and not sottotitolo:
            sottotitolo = riga[4:].strip()
            sottotitolo_riga = riga.strip()
        if riga.startswith('Ultimo aggiornamento'):
            data = riga.strip()
    copertina(doc, sottotitolo, data)

    i = 0
    while i < len(testo):
        riga = testo[i].rstrip()

        # titolo, sottotitolo e data stanno gia' in copertina
        if riga.startswith('# ') or riga.strip() == sottotitolo_riga or riga.startswith('Ultimo aggiornamento'):
            i += 1
            continue

        if not riga.strip():
            i += 1
            continue

        if riga.startswith('---'):
            i += 1
            continue

        # tabelle
        if riga.startswith('|') and i + 1 < len(testo) and re.match(r'^\|[\s:\-|]+\|$', testo[i + 1].strip()):
            blocco = []
            while i < len(testo) and testo[i].strip().startswith('|'):
                blocco.append(testo[i].strip())
                i += 1
            tabella(doc, blocco)
            continue

        # citazioni: si uniscono le righe consecutive
        if riga.startswith('>'):
            pezzi = []
            while i < len(testo) and testo[i].strip().startswith('>'):
                pezzi.append(testo[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(10)
            bordo_sinistro(p, '7A3FD6')
            scrivi(p, ' '.join(x for x in pezzi if x))
            for r in p.runs:
                r.font.size = Pt(10.5)
            continue

        # titoli
        liv = len(riga) - len(riga.lstrip('#'))
        if 0 < liv <= 4 and riga[liv:liv + 1] == ' ':
            testo_tit = re.sub(r'\*+', '', riga[liv + 1:].strip())
            doc.add_heading(testo_tit, min(liv, 4) - 1 if liv > 1 else 1)
            i += 1
            continue

        # elenchi
        m = re.match(r'^(\d+)\.\s+(.*)$', riga)
        if m:
            p = doc.add_paragraph(style='List Number')
            scrivi(p, m.group(2))
            i += 1
            while i < len(testo) and testo[i].startswith('   ') and testo[i].strip():
                scrivi(p, ' ' + testo[i].strip())
                i += 1
            continue
        if riga.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            scrivi(p, riga[2:])
            i += 1
            while i < len(testo) and testo[i].startswith('  ') and testo[i].strip():
                scrivi(p, ' ' + testo[i].strip())
                i += 1
            continue

        # paragrafo normale: le righe spezzate del markdown si ricuciono
        pezzi = []
        while i < len(testo) and testo[i].strip() and not re.match(
                r'^(#{1,4} |[-*] |\d+\. |\||>|---)', testo[i]):
            pezzi.append(testo[i].strip())
            i += 1
        p = doc.add_paragraph()
        scrivi(p, ' '.join(pezzi))

    doc.save(USCITA)
    return USCITA


if __name__ == '__main__':
    print(converti())
